from docx import Document
import os
import uuid
import datetime
import logging
from django.conf import settings

logger = logging.getLogger(__name__)

def generate_docx(data):
    """Gera DOCX a partir do template com os dados fornecidos."""
    logger.info("Iniciando geração do DOCX")
    
    # Caminho do template
    template_path = os.path.join(settings.BASE_DIR, 'merito_academico', 'templates_documents', 'Carta.docx')
    if not os.path.exists(template_path):
        logger.error(f"Template não encontrado: {template_path}")
        raise FileNotFoundError("Template não encontrado")

    # Carrega o documento
    doc = Document(template_path)

    # Prepara os dados para substituição
    replacements = {
        '[Nome do Diretor ou Diretora]': data['school_Principal_Name'],
        '[Nome da Escola]': data['school_name'],
        '[Endereço Completo da Escola]': data['school_complete_address'],
        '[Cidade]': 'Itabira',
        '[Data]': datetime.datetime.now().strftime('%d/%m/%Y'),
        '[Nome do Aluno]': f"{data['first_name']} {data['last_name']}",
        '[Nota]': str(data['HighSchool_score']),
    }

    # Substitui os placeholders em todos os parágrafos
    for paragraph in doc.paragraphs:
        for placeholder, value in replacements.items():
            if placeholder in paragraph.text:
                paragraph.text = paragraph.text.replace(placeholder, value)

    # Salva o DOCX temporário com UUID
    temp_docx = f'temp_{uuid.uuid4()}.docx'
    temp_docx_path = os.path.join(settings.MEDIA_ROOT, temp_docx)
    doc.save(temp_docx_path)
    logger.info(f"DOCX salvo: {temp_docx_path}")

    return temp_docx_path